"""Tools the agent can call. Each tool takes (state, **args) and returns a string observation."""
import os

from pptx import Presentation
from pptx.util import Pt

from .best_practices import MAX_BULLETS, MAX_TITLE_WORDS, MAX_WORDS_PER_BULLET


class DeckState:
    def __init__(self, source_text: str = ""):
        self.source_text = source_text
        self.outline: list[str] = []
        self.slides: list[dict] = []  # {"layout", "title", "bullets", "notes"}
        self.saved_path: str | None = None


def read_source(state, start: int = 0, length: int = 4000):
    if not state.source_text:
        return "No source document was provided."
    chunk = state.source_text[start:start + length]
    return f"[chars {start}-{start + len(chunk)} of {len(state.source_text)}]\n{chunk}"


def plan_outline(state, titles: list):
    state.outline = [str(t) for t in titles]
    return f"Outline saved with {len(state.outline)} slides."


def add_slide(state, title: str, bullets: list = None, notes: str = "", layout: str = "content"):
    if layout not in ("title", "content", "section"):
        layout = "content"
    state.slides.append({
        "layout": layout,
        "title": str(title),
        "bullets": [str(b) for b in (bullets or [])],
        "notes": str(notes),
    })
    return f"Added slide {len(state.slides)}: {title}"


def edit_slide(state, index: int, title: str = None, bullets: list = None, notes: str = None):
    i = int(index) - 1  # 1-based for the LLM
    if not 0 <= i < len(state.slides):
        return f"Error: no slide {index}. Deck has {len(state.slides)} slides."
    s = state.slides[i]
    if title is not None:
        s["title"] = str(title)
    if bullets is not None:
        s["bullets"] = [str(b) for b in bullets]
    if notes is not None:
        s["notes"] = str(notes)
    return f"Edited slide {index}."


def delete_slide(state, index: int):
    i = int(index) - 1
    if not 0 <= i < len(state.slides):
        return f"Error: no slide {index}."
    removed = state.slides.pop(i)
    return f"Deleted slide {index}: {removed['title']}"


def view_deck(state):
    if not state.slides:
        return "Deck is empty."
    lines = []
    for n, s in enumerate(state.slides, 1):
        lines.append(f"{n}. [{s['layout']}] {s['title']}")
        for b in s["bullets"]:
            lines.append(f"   - {b}")
    return "\n".join(lines)


def critique_deck(state):
    """Rule-based check against the best practices."""
    issues = []
    if len(state.slides) < 3:
        issues.append("Deck is very short; most decks need a title, content, and a summary slide.")
    if state.slides and state.slides[0]["layout"] != "title":
        issues.append("Slide 1 should use the 'title' layout.")
    last_titles = " ".join(s["title"].lower() for s in state.slides[-2:])
    if state.slides and not any(k in last_titles for k in ("summary", "takeaway", "next step", "question", "conclusion", "q&a")):
        issues.append("Deck should end with a summary/takeaways or next steps slide.")
    for n, s in enumerate(state.slides, 1):
        if len(s["title"].split()) > MAX_TITLE_WORDS:
            issues.append(f"Slide {n}: title too long ({len(s['title'].split())} words).")
        if s["layout"] == "content":
            if not s["bullets"]:
                issues.append(f"Slide {n}: content slide has no bullets.")
            if len(s["bullets"]) > MAX_BULLETS:
                issues.append(f"Slide {n}: {len(s['bullets'])} bullets (max {MAX_BULLETS}).")
            for b in s["bullets"]:
                if len(b.split()) > MAX_WORDS_PER_BULLET:
                    issues.append(f"Slide {n}: bullet too long: '{b[:40]}...'")
            if not s["notes"]:
                issues.append(f"Slide {n}: missing speaker notes.")
    return "No issues found." if not issues else "Issues:\n" + "\n".join(f"- {i}" for i in issues)


def save_deck(state, filename: str = "deck.pptx", output_dir: str = "output"):
    os.makedirs(output_dir, exist_ok=True)
    if not filename.endswith(".pptx"):
        filename += ".pptx"
    path = os.path.join(output_dir, os.path.basename(filename))

    prs = Presentation()
    for s in state.slides:
        if s["layout"] == "title":
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = s["title"]
            if s["bullets"]:
                slide.placeholders[1].text = s["bullets"][0]
        elif s["layout"] == "section":
            slide = prs.slides.add_slide(prs.slide_layouts[2])
            slide.shapes.title.text = s["title"]
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = s["title"]
            tf = slide.placeholders[1].text_frame
            for j, b in enumerate(s["bullets"]):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = b
                p.font.size = Pt(24)
        if s["notes"]:
            slide.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(path)
    state.saved_path = path
    return f"Saved deck to {path}"


TOOLS = {
    "read_source": read_source,
    "plan_outline": plan_outline,
    "add_slide": add_slide,
    "edit_slide": edit_slide,
    "delete_slide": delete_slide,
    "view_deck": view_deck,
    "critique_deck": critique_deck,
    "save_deck": save_deck,
}

TOOL_DOCS = """\
- read_source(start:int=0, length:int=4000): read part of the source document.
- plan_outline(titles:list[str]): record the planned slide titles.
- add_slide(title:str, bullets:list[str], notes:str, layout:"title"|"content"|"section"): append a slide. For a "title" slide, bullets[0] is the subtitle.
- edit_slide(index:int, title?:str, bullets?:list[str], notes?:str): modify a slide (1-based).
- delete_slide(index:int): remove a slide (1-based).
- view_deck(): list the current slides.
- critique_deck(): check the deck against the best practices.
- save_deck(filename:str): write the .pptx file. Call this when done.
- finish(message:str): end the task after saving.
"""
